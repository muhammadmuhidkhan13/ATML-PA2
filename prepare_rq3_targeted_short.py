"""Targeted, blinded RQ3 qualitative comparison; uses openpyxl on the user's PC.

Run: python prepare_rq3_targeted_short.py --per-tier 2 --seed 6304
Requires: python -m pip install openpyxl
"""
import argparse
import csv
import json
import random
from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SOURCES = {
    'beta_0': 'kl_0_eval_generations.jsonl',
    'beta_0p1': 'kl_0p1_eval_generations.jsonl',
    'beta_0p2': 'kl_0p2_eval_generations.jsonl',
}


def load_jsonl(path):
    if not path.is_file():
        raise FileNotFoundError(f'Missing: {path}')
    with path.open(encoding='utf-8') as handle:
        rows = [json.loads(line) for line in handle if line.strip()]
    result = {str(r['evaluation_index']): r for r in rows}
    if len(result) != len(rows):
        raise ValueError(f'Duplicate evaluation_index in {path}')
    return result


def render_prompt(messages):
    if isinstance(messages, str):
        return messages
    return '\n\n'.join(f"{m.get('role', 'unknown').upper()}: {m.get('content', '')}" for m in messages)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--results-dir', default='results/task2_ppo')
    parser.add_argument('--per-tier', type=int, default=2, help='Cases each for low/middle/high; default 2 = 6 prompts and 18 responses')
    parser.add_argument('--max-prompt-chars', type=int, default=1100)
    parser.add_argument('--max-response-chars', type=int, default=1800)
    parser.add_argument('--seed', type=int, default=6304)
    args = parser.parse_args()
    if args.per_tier < 1:
        parser.error('--per-tier must be >= 1')
    root = Path(args.results_dir)
    tables = {k: load_jsonl(root / path) for k, path in SOURCES.items()}
    common = set.intersection(*(set(t) for t in tables.values()))
    candidates = []
    for idx in common:
        rows = [tables[k][idx] for k in SOURCES]
        if len({str(r.get('prompt_id')) for r in rows}) != 1:
            continue
        if any(r.get('reward_model_score') is None for r in rows):
            continue
        prompt = render_prompt(rows[0].get('prompt_messages', []))
        if len(prompt) > args.max_prompt_chars:
            continue
        if any(len(str(r.get('response') or '')) > args.max_response_chars for r in rows):
            continue
        # Stratify on mean learned reward across matched responses, not on any one policy.
        candidates.append((float(mean(float(r['reward_model_score']) for r in rows)), int(idx), prompt))
    candidates.sort()
    n = len(candidates)
    need = 3 * args.per_tier
    if n < need:
        raise ValueError(f'Only {n} short eligible matched prompts, need {need}. '
                         'Try --max-prompt-chars 1800 --max-response-chars 3000.')

    # Disjoint tertiles of the eligible short-prompt pool.
    q1, q2 = n // 3, 2 * n // 3
    buckets = {'Low': candidates[:q1], 'Middle': candidates[q1:q2], 'High': candidates[q2:]}
    rng = random.Random(args.seed)
    selected = []
    for tier, bucket in buckets.items():
        if len(bucket) < args.per_tier:
            raise ValueError(f'Only {len(bucket)} eligible {tier} cases; need {args.per_tier}.')
        selected.extend((tier, item) for item in rng.sample(bucket, args.per_tier))
    # Keep tier order; random within each tier for reproducible blind review.

    wb = Workbook()
    ws = wb.active
    ws.title = 'Blind Review'
    headers = ['Case', 'Prompt', 'Response', 'Generated answer', 'Overall quality (1–5)',
               'Instruction following (1–5)', 'Factual accuracy (1–5 or N/A)',
               'Excess verbosity (1 none–5 severe)', 'Notes']
    ws.append(headers)
    reveal = []
    shuffle_rng = random.Random(args.seed + 1)
    for case, (tier, (mean_reward, idx, prompt)) in enumerate(selected, 1):
        order = list(SOURCES)
        shuffle_rng.shuffle(order)
        for letter, policy in zip('ABC', order):
            row = tables[policy][str(idx)]
            ws.append([case, prompt, letter, str(row.get('response') or ''), None, None, None, None, None])
            reveal.append({
                'case': case, 'evaluation_index': idx, 'anonymous_response': letter,
                'policy': policy, 'reward_tier': tier,
                'reward_model_score': row['reward_model_score'],
                'response_length': row.get('response_length'),
                'sampled_signed_kl': row.get('sampled_kl'),
                'terminated_with_eos': row.get('terminated_with_eos'),
                'truncated': row.get('truncated'),
            })

    widths = {'A': 9, 'B': 73, 'C': 13, 'D': 90, 'E': 23,
              'F': 25, 'G': 26, 'H': 28, 'I': 49}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17324D')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 36
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 145
        for col in range(1, 10):
            c = ws.cell(row, col)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if col >= 5:
                c.fill = PatternFill('solid', fgColor='EAF4EC')
            elif (row - 2) // 3 % 2:
                c.fill = PatternFill('solid', fgColor='F2F5F8')
    dv = DataValidation(type='list', formula1='"1,2,3,4,5"', allow_blank=True)
    dv.error = 'Choose an integer from 1 to 5.'
    dv.errorTitle = 'Invalid rating'
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    for col in ['E', 'F', 'H']:
        dv.add(f'{col}2:{col}{ws.max_row}')
    dv_acc = DataValidation(type='list', formula1='"1,2,3,4,5,N/A"', allow_blank=True)
    ws.add_data_validation(dv_acc)
    dv_acc.add(f'G2:G{ws.max_row}')
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:I{ws.max_row}'

    instructions = wb.create_sheet('Instructions')
    notes = [
        'RQ3 qualitative review — blind ratings',
        'This is a TARGETED, STRATIFIED CASE SAMPLE, not a random representative sample.',
        'The selection uses low, middle, and high mean learned reward across matched model outputs.',
        'Reward strata and actual scores are deliberately NOT included in this workbook.',
        'Rate all 3 responses to each prompt independently, before opening the reward-key CSV.',
        'Higher overall quality, instruction following, and factual accuracy = better.',
        'For verbosity: 1 = no unnecessary verbosity; 5 = severely verbose/repetitive (WORSE).',
        'Use N/A for factual accuracy when the prompt is subjective or cannot be verified.',
        'Avoid claims of overall reward-model accuracy or reward overoptimization from this selected sample alone.',
        'A disagreeing response is a potential reward-model limitation, not by itself proof of PPO reward hacking.',
        'Do not disclose or open the reward key until after scoring.',
    ]
    for i, note in enumerate(notes, 1):
        cell = instructions.cell(i, 1, note)
        cell.alignment = Alignment(wrap_text=True)
        if i == 1:
            cell.font = Font(size=15, bold=True, color='17324D')
    instructions.column_dimensions['A'].width = 120
    for i in range(1, len(notes) + 1):
        instructions.row_dimensions[i].height = 32

    review_file = root / 'rq3_targeted_short_blind_review.xlsx'
    key_file = root / 'rq3_targeted_short_REWARD_KEY_DO_NOT_OPEN.csv'
    root.mkdir(parents=True, exist_ok=True)
    wb.save(review_file)
    with key_file.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(reveal[0]))
        writer.writeheader()
        writer.writerows(reveal)

    print(f'Matched short-prompt candidates: {n}')
    print(f'Selected: {len(selected)} prompts, {len(reveal)} responses')
    print(f'Reward tiers: {args.per_tier} low / {args.per_tier} middle / {args.per_tier} high')
    print('BLIND workbook: ', review_file)
    print('SEALED reward key:', key_file)
    print('Only open the reward key AFTER all quality ratings are entered.')

if __name__ == '__main__':
    main()
