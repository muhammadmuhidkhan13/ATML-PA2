from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


POLICY_FILES = {
    "kl_0": "kl_0_eval_generations.jsonl",
    "kl_0p1": "kl_0p1_eval_generations.jsonl",
    "kl_0p2": "kl_0p2_eval_generations.jsonl",
}

REQUIRED_FIELDS = {
    "evaluation_index",
    "prompt_id",
    "prompt_messages",
    "response",
    "reward_model_score",
    "sampled_kl",
    "response_length",
    "terminated_with_eos",
    "truncated",
}

KEY_COLUMNS = [
    "review_item",
    "prompt_group",
    "selection_order",
    "evaluation_index",
    "prompt_id",
    "anonymous_label",
    "actual_policy",
    "source_file",
    "reward_model_score",
    "sampled_kl",
    "response_length",
    "terminated_with_eos",
    "truncated",
    "response_sha256",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create a blind Excel review workbook and a separate reward key "
            "from matched PPO KL-ablation generations."
        )
    )
    parser.add_argument(
        "--results-dir",
        type=Path,
        default=Path("results/task2_ppo"),
    )
    parser.add_argument("--number", type=int, default=15)
    parser.add_argument("--seed", type=int, default=6304)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=None,
        help="Default: RESULTS_DIR/rq3_blind_quality_review.xlsx",
    )
    parser.add_argument(
        "--key",
        type=Path,
        default=None,
        help=(
            "Default: RESULTS_DIR/"
            "rq3_reward_key_DO_NOT_OPEN_UNTIL_RATED.csv"
        ),
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Replace only the two generated review outputs if they exist.",
    )
    return parser.parse_args()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise FileNotFoundError(f"Missing evaluation file: {path}")

    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Invalid JSON in {path} at line {line_number}: {exc}"
                ) from exc

            missing = REQUIRED_FIELDS.difference(record)
            if missing:
                raise ValueError(
                    f"{path}, line {line_number}, is missing fields: "
                    f"{sorted(missing)}"
                )
            records.append(record)

    if not records:
        raise ValueError(f"Evaluation file is empty: {path}")
    return records


def index_records(
    records: list[dict[str, Any]],
    path: Path,
) -> dict[int, dict[str, Any]]:
    indexed: dict[int, dict[str, Any]] = {}
    for record in records:
        try:
            evaluation_index = int(record["evaluation_index"])
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"Non-integer evaluation_index in {path}: "
                f"{record['evaluation_index']!r}"
            ) from exc
        if evaluation_index in indexed:
            raise ValueError(
                f"Duplicate evaluation_index {evaluation_index} in {path}"
            )
        indexed[evaluation_index] = record
    return indexed


def canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def render_prompt(messages: Any) -> str:
    if isinstance(messages, str):
        return messages.strip()
    if not isinstance(messages, list):
        return canonical_json(messages)

    rendered: list[str] = []
    for message in messages:
        if isinstance(message, dict):
            role = str(message.get("role", "message")).upper()
            content = message.get("content", "")
            if isinstance(content, str):
                content_text = content
            else:
                content_text = canonical_json(content)
            rendered.append(f"{role}: {content_text}")
        else:
            rendered.append(str(message))
    return "\n\n".join(rendered).strip()


def excel_safe_text(value: Any) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def text_sha256(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_and_match(
    indexed: dict[str, dict[int, dict[str, Any]]],
) -> list[int]:
    policies = list(POLICY_FILES)
    expected_indices = set(indexed[policies[0]])
    for policy in policies[1:]:
        actual_indices = set(indexed[policy])
        if actual_indices != expected_indices:
            missing = sorted(expected_indices.difference(actual_indices))
            extra = sorted(actual_indices.difference(expected_indices))
            raise ValueError(
                f"Evaluation-index set mismatch for {policy}. "
                f"Missing={missing[:10]}, extra={extra[:10]}"
            )

    if not expected_indices:
        raise ValueError("The three evaluation files have no matched indices.")

    for evaluation_index in sorted(expected_indices):
        rows = [indexed[p][evaluation_index] for p in policies]
        prompt_ids = {canonical_json(row["prompt_id"]) for row in rows}
        if len(prompt_ids) != 1:
            raise ValueError(
                "prompt_id mismatch at evaluation_index "
                f"{evaluation_index}: {[row['prompt_id'] for row in rows]}"
            )

        prompts = {
            canonical_json(row["prompt_messages"])
            for row in rows
        }
        if len(prompts) != 1:
            raise ValueError(
                "prompt_messages mismatch at evaluation_index "
                f"{evaluation_index}"
            )
    return sorted(expected_indices)


def create_blind_rows(
    indexed: dict[str, dict[int, dict[str, Any]]],
    matched_indices: list[int],
    number: int,
    seed: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[int]]:
    if number < 1:
        raise ValueError("--number must be at least 1.")
    if number > len(matched_indices):
        raise ValueError(
            f"Requested {number} prompts but only "
            f"{len(matched_indices)} matched prompts are available."
        )

    selection_rng = random.Random(seed)
    selected = selection_rng.sample(matched_indices, number)

    # Keep assignment randomness independent from the sampling call count.
    assignment_rng = random.Random(seed + 1)
    policies = list(POLICY_FILES)
    anonymous_labels = ["A", "B", "C"]

    blind_rows: list[dict[str, Any]] = []
    key_rows: list[dict[str, Any]] = []

    for selection_order, evaluation_index in enumerate(selected, start=1):
        shuffled_policies = policies.copy()
        assignment_rng.shuffle(shuffled_policies)
        prompt_group = f"P{selection_order:02d}"

        for label, policy in zip(anonymous_labels, shuffled_policies):
            source = indexed[policy][evaluation_index]
            review_item = f"{prompt_group}-{label}"
            response = str(source["response"])

            blind_rows.append(
                {
                    "review_item": review_item,
                    "prompt_group": prompt_group,
                    "prompt": render_prompt(source["prompt_messages"]),
                    "anonymous_response": response,
                }
            )
            key_rows.append(
                {
                    "review_item": review_item,
                    "prompt_group": prompt_group,
                    "selection_order": selection_order,
                    "evaluation_index": evaluation_index,
                    "prompt_id": source["prompt_id"],
                    "anonymous_label": label,
                    "actual_policy": policy,
                    "source_file": POLICY_FILES[policy],
                    "reward_model_score": source["reward_model_score"],
                    "sampled_kl": source["sampled_kl"],
                    "response_length": source["response_length"],
                    "terminated_with_eos": source["terminated_with_eos"],
                    "truncated": source["truncated"],
                    "response_sha256": text_sha256(response),
                }
            )
    return blind_rows, key_rows, selected


def build_workbook(
    blind_rows: list[dict[str, Any]],
    output_path: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blind Review"
    sheet.sheet_view.showGridLines = False

    navy = "17365D"
    blue = "D9EAF7"
    pale_blue = "EDF4FA"
    pale_gray = "F3F4F6"
    white = "FFFFFF"
    rating_fill = "FFF2CC"
    red_fill = "FCE4D6"
    thin_gray = Side(style="thin", color="B7C3D0")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    sheet.merge_cells("A1:I1")
    sheet["A1"] = "Task 2 RQ3 — Blind Response Quality Review"
    sheet["A1"].font = Font(size=16, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells("A2:I2")
    sheet["A2"] = (
        "Rate each response independently without trying to infer its source. "
        "Use whole numbers from 1 (worst) to 5 (best), except verbosity, "
        "where 1 = concise/no problem and 5 = severely verbose or repetitive."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A2"].fill = PatternFill("solid", fgColor=blue)
    sheet.row_dimensions[2].height = 44

    sheet.merge_cells("A3:I3")
    sheet["A3"] = (
        "Factual accuracy: choose N/A when the response contains no meaningful "
        "factual claim or the prompt does not permit factual assessment. "
        "Use Notes for specific errors, omissions, instruction violations, "
        "or reasons for unusually high/low ratings."
    )
    sheet["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A3"].fill = PatternFill("solid", fgColor=pale_blue)
    sheet.row_dimensions[3].height = 44

    sheet.merge_cells("A4:I4")
    sheet["A4"] = (
        "Do not open the separate reward-key CSV until every rating is complete "
        "and this workbook has been saved. Yellow cells are editable ratings."
    )
    sheet["A4"].font = Font(bold=True, color="9C0006")
    sheet["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet["A4"].fill = PatternFill("solid", fgColor=red_fill)
    sheet.row_dimensions[4].height = 32

    headers = [
        "Review item",
        "Prompt group",
        "Original prompt",
        "Anonymous response",
        "Overall quality\n(1 worst–5 best)",
        "Instruction following\n(1 worst–5 best)",
        "Factual accuracy\n(1 worst–5 best or N/A)",
        "Unnecessary verbosity/repetition\n(1 none–5 severe; higher is worse)",
        "Notes",
    ]
    header_row = 6
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    sheet.row_dimensions[header_row].height = 52

    for offset, record in enumerate(blind_rows, start=1):
        row_number = header_row + offset
        group_number = (offset - 1) // 3
        base_fill = blue if group_number % 2 == 0 else pale_gray
        values = [
            record["review_item"],
            record["prompt_group"],
            excel_safe_text(record["prompt"]),
            excel_safe_text(record["anonymous_response"]),
            "",
            "",
            "",
            "",
            "",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            cell.fill = PatternFill(
                "solid",
                fgColor=rating_fill if column >= 5 else base_fill,
            )
        sheet.cell(row_number, 1).font = Font(bold=True)
        sheet.cell(row_number, 2).font = Font(bold=True)
        sheet.row_dimensions[row_number].height = 150

    last_row = header_row + len(blind_rows)
    rating_validation = DataValidation(
        type="list",
        formula1='"1,2,3,4,5"',
        allow_blank=True,
    )
    rating_validation.error = "Choose a whole-number rating from 1 to 5."
    rating_validation.errorTitle = "Invalid rating"
    rating_validation.prompt = "Choose 1, 2, 3, 4, or 5."
    rating_validation.promptTitle = "Rating scale"
    rating_validation.showErrorMessage = True
    rating_validation.showInputMessage = True
    sheet.add_data_validation(rating_validation)
    for column in ("E", "F", "H"):
        rating_validation.add(f"{column}{header_row + 1}:{column}{last_row}")

    factual_validation = DataValidation(
        type="list",
        formula1='"1,2,3,4,5,N/A"',
        allow_blank=True,
    )
    factual_validation.error = "Choose 1–5 or N/A."
    factual_validation.errorTitle = "Invalid rating"
    factual_validation.showErrorMessage = True
    factual_validation.showInputMessage = True
    sheet.add_data_validation(factual_validation)
    factual_validation.add(f"G{header_row + 1}:G{last_row}")

    widths = {
        "A": 14,
        "B": 14,
        "C": 48,
        "D": 72,
        "E": 17,
        "F": 19,
        "G": 19,
        "H": 24,
        "I": 42,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "C7"
    sheet.auto_filter.ref = f"A{header_row}:I{last_row}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.auto_filter.ref = f"A{header_row}:I{last_row}"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)


def write_key(key_rows: list[dict[str, Any]], output_path: Path) -> None:
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=KEY_COLUMNS)
        writer.writeheader()
        for record in key_rows:
            row = dict(record)
            row["prompt_id"] = canonical_json(row["prompt_id"])
            writer.writerow(row)


def verify_outputs(
    workbook_path: Path,
    key_path: Path,
    blind_rows: list[dict[str, Any]],
    key_rows: list[dict[str, Any]],
) -> None:
    expected_count = len(blind_rows)
    workbook = load_workbook(workbook_path, data_only=False)
    if workbook.sheetnames != ["Blind Review"]:
        raise RuntimeError(
            f"Workbook must contain exactly one visible sheet; found "
            f"{workbook.sheetnames}"
        )
    sheet = workbook["Blind Review"]
    if sheet.sheet_state != "visible":
        raise RuntimeError("Blind Review worksheet is not visible.")

    workbook_rows: dict[str, str] = {}
    for row_number in range(7, 7 + expected_count):
        review_item = sheet.cell(row_number, 1).value
        response = sheet.cell(row_number, 4).value
        if not review_item or response is None:
            raise RuntimeError(f"Incomplete workbook row {row_number}.")
        response = str(response)
        if response.startswith("'") and response[1:].startswith(
            ("=", "+", "-", "@")
        ):
            response = response[1:]
        workbook_rows[str(review_item)] = response

    for row in sheet.iter_rows():
        for cell in row:
            if cell.comment is not None:
                raise RuntimeError(
                    f"Unexpected workbook comment in cell {cell.coordinate}."
                )

    if len(workbook_rows) != expected_count:
        raise RuntimeError("Workbook review-item IDs are not unique.")

    with key_path.open("r", encoding="utf-8-sig", newline="") as handle:
        parsed_key = list(csv.DictReader(handle))
    if len(parsed_key) != expected_count:
        raise RuntimeError(
            f"Expected {expected_count} key rows; found {len(parsed_key)}."
        )

    expected_by_item = {row["review_item"]: row for row in key_rows}
    actual_by_item = {row["review_item"]: row for row in parsed_key}
    if set(expected_by_item) != set(actual_by_item):
        raise RuntimeError("Workbook/key review-item IDs do not align.")

    for review_item, response in workbook_rows.items():
        expected_hash = expected_by_item[review_item]["response_sha256"]
        if text_sha256(response) != expected_hash:
            raise RuntimeError(
                f"Response hash mismatch for review item {review_item}."
            )

    group_to_labels: dict[str, list[str]] = {}
    group_to_policies: dict[str, list[str]] = {}
    for row in key_rows:
        group_to_labels.setdefault(row["prompt_group"], []).append(
            row["anonymous_label"]
        )
        group_to_policies.setdefault(row["prompt_group"], []).append(
            row["actual_policy"]
        )

    expected_labels = sorted(["A", "B", "C"])
    expected_policies = sorted(POLICY_FILES)
    for group in group_to_labels:
        if sorted(group_to_labels[group]) != expected_labels:
            raise RuntimeError(f"Invalid anonymous labels in {group}.")
        if sorted(group_to_policies[group]) != expected_policies:
            raise RuntimeError(f"Invalid policy mapping in {group}.")

    # Inspect all ordinary cell values for direct leakage of policy/file names.
    forbidden = set(POLICY_FILES)
    forbidden.update(POLICY_FILES.values())
    for row in sheet.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            if any(token in value for token in forbidden):
                raise RuntimeError(
                    f"Policy identity leaked into workbook cell {cell.coordinate}."
                )

    # Also scan the complete XLSX package, including document properties.
    with zipfile.ZipFile(workbook_path, "r") as archive:
        for member in archive.namelist():
            payload = archive.read(member)
            for token in forbidden:
                if token.encode("utf-8") in payload:
                    raise RuntimeError(
                        "Policy identity leaked into workbook package member "
                        f"{member}."
                    )


def main() -> None:
    args = parse_args()
    results_dir = args.results_dir.resolve()
    workbook_path = (
        args.workbook
        if args.workbook is not None
        else results_dir / "rq3_blind_quality_review.xlsx"
    ).resolve()
    key_path = (
        args.key
        if args.key is not None
        else results_dir
        / "rq3_reward_key_DO_NOT_OPEN_UNTIL_RATED.csv"
    ).resolve()

    if workbook_path == key_path:
        raise ValueError("Workbook and key paths must be different.")
    for output_path in (workbook_path, key_path):
        if output_path.exists() and not args.overwrite:
            raise FileExistsError(
                f"Refusing to overwrite existing output: {output_path}. "
                "Move it, rename it, or rerun with --overwrite."
            )
        output_path.parent.mkdir(parents=True, exist_ok=True)

    indexed: dict[str, dict[int, dict[str, Any]]] = {}
    source_counts: dict[str, int] = {}
    for policy, filename in POLICY_FILES.items():
        path = results_dir / filename
        records = read_jsonl(path)
        indexed[policy] = index_records(records, path)
        source_counts[policy] = len(records)

    matched_indices = validate_and_match(indexed)
    blind_rows, key_rows, selected = create_blind_rows(
        indexed=indexed,
        matched_indices=matched_indices,
        number=args.number,
        seed=args.seed,
    )

    build_workbook(blind_rows, workbook_path)
    write_key(key_rows, key_path)
    verify_outputs(workbook_path, key_path, blind_rows, key_rows)

    policy_counts = Counter(row["actual_policy"] for row in key_rows)
    print("\nRQ3 BLIND QUALITY REVIEW CREATED")
    print("=" * 84)
    print(f"Source records: {source_counts}")
    print(f"Matched prompt indices available: {len(matched_indices)}")
    print(f"Selection seed: {args.seed}")
    print(f"Anonymization seed: {args.seed + 1}")
    print(f"Selected evaluation indices: {selected}")
    print(f"Selected prompts: {len(selected)}")
    print(f"Blind response rows: {len(blind_rows)}")
    print(f"Rows per policy in sealed key: {dict(policy_counts)}")
    print("Prompt-ID consistency: PASS")
    print("Prompt-text consistency: PASS")
    print("Each prompt contains A, B, and C exactly once: PASS")
    print("Each prompt contains all three policies exactly once: PASS")
    print("Workbook/key response-hash alignment: PASS")
    print("Workbook identity/reward leakage check: PASS")
    print(f"\nWorkbook: {workbook_path}")
    print(f"Workbook SHA-256: {file_sha256(workbook_path)}")
    print(f"Reward key: {key_path}")
    print(f"Reward-key SHA-256: {file_sha256(key_path)}")
    print("\nDo not open the reward-key CSV until all ratings are complete.")


if __name__ == "__main__":
    main()
